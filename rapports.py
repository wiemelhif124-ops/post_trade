import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal
from xml.sax.saxutils import escape

import pandas as pd
from sqlalchemy import text

ROUGE = "#E8615B"


# ---------------------------------------------------------------- utilitaires
def fmt(v):
    """Nombre lisible : 1 800 000 ou 12,50."""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    if f != f:
        return "—"
    if f == int(f):
        return f"{f:,.0f}".replace(",", " ")
    return f"{f:,.2f}".replace(",", " ").replace(".", ",")


def _lire(engine, sql, **params):
    with engine.connect() as conn:
        df = pd.read_sql_query(text(sql), conn, params=params or None)
    for col in df.columns:  # PostgreSQL/Oracle renvoient des Decimal
        if df[col].dtype == object and len(df) and isinstance(df[col].iloc[0], Decimal):
            df[col] = df[col].astype(float)
    return df


def _est_nombre(df, c):
    return pd.api.types.is_numeric_dtype(df[c]) and not pd.api.types.is_bool_dtype(df[c])


def _texte(v, court=False):
    if v is None or (isinstance(v, float) and v != v):
        return ""
    if isinstance(v, (date, datetime)):
        return v.strftime("%d/%m" if court else "%d/%m/%Y")
    return str(v)


def _section(titre, df, vue="tableau"):
    return {"titre": titre, "df": df, "vue": vue}


# ---------------------------------------------------------------- les rapports
def _rapport_reglement(engine, d1, d2):
    p = dict(d1=d1, d2=d2)
    entre = "i.date_reglement BETWEEN :d1 AND :d2"
    statuts = _lire(engine, f"SELECT i.statut, COUNT(*) AS nb FROM instructions i WHERE {entre} "
                            "GROUP BY i.statut ORDER BY nb DESC", **p)
    nb = {r.statut: int(r.nb) for r in statuts.itertuples()}
    total = sum(nb.values())
    reglees, echouees = nb.get("regle", 0), nb.get("echoue", 0)
    taux = round(100 * reglees / (reglees + echouees), 1) if (reglees + echouees) else "n/a"
    indicateurs = [("Instructions", total), ("Réglées", reglees), ("Échouées", echouees),
                   ("En attente ou appariées", nb.get("en_attente", 0) + nb.get("apparie", 0)),
                   ("Taux de règlement (%)", taux)]
    sections = [
        _section("Instructions par statut", statuts, "histogramme"),
        _section("Échecs par motif", _lire(
            engine, f"SELECT COALESCE(i.motif_echec, 'Non précisé') AS motif, COUNT(*) AS nb FROM instructions i "
                    f"WHERE i.statut='echoue' AND {entre} GROUP BY i.motif_echec ORDER BY nb DESC", **p), "histogramme"),
        _section("Règlements par jour", _lire(
            engine, f"SELECT i.date_reglement, COUNT(*) AS nb FROM instructions i WHERE {entre} "
                    "GROUP BY i.date_reglement ORDER BY i.date_reglement", **p), "courbe"),
        _section("Volumes par devise", _lire(
            engine, f"SELECT i.devise, COUNT(*) AS nb_instructions, SUM(i.montant) AS montant_total "
                    f"FROM instructions i WHERE {entre} GROUP BY i.devise ORDER BY i.devise", **p)),
        _section("Détail des instructions échouées", _lire(
            engine, "SELECT i.id_instruction, c.nom AS client, t.nom AS titre, i.sens, i.quantite, i.montant, "
                    "i.devise, i.date_reglement, i.motif_echec FROM instructions i "
                    "JOIN comptes_clients cc ON cc.id_compte = i.id_compte "
                    "JOIN clients c ON c.id_client = cc.id_client JOIN titres t ON t.isin = i.isin "
                    f"WHERE i.statut='echoue' AND {entre} ORDER BY i.date_reglement DESC", **p)),
    ]
    return indicateurs, sections


def _rapport_a_venir(engine, d1, d2):
    p = dict(d1=d1, d2=d2)
    ouvert = "i.statut IN ('en_attente','apparie') AND i.date_reglement BETWEEN :d1 AND :d2"
    statuts = _lire(engine, f"SELECT i.statut, COUNT(*) AS nb FROM instructions i WHERE {ouvert} "
                            "GROUP BY i.statut", **p)
    nb = {r.statut: int(r.nb) for r in statuts.itertuples()}
    risque = _lire(
        engine, "SELECT i.id_instruction, c.nom AS client, i.montant AS montant_achat, i.devise, "
                "cash.solde AS solde_cash, i.date_reglement FROM instructions i "
                "JOIN comptes_clients cc ON cc.id_compte = i.id_compte "
                "JOIN clients c ON c.id_client = cc.id_client "
                "JOIN comptes_clients cash ON cash.id_client = cc.id_client "
                "AND cash.type_compte = 'cash' AND cash.devise = i.devise "
                f"WHERE i.sens = 'achat' AND {ouvert} AND cash.solde < i.montant ORDER BY i.date_reglement", **p)
    indicateurs = [("Règlements à venir", sum(nb.values())), ("En attente", nb.get("en_attente", 0)),
                   ("Appariés", nb.get("apparie", 0)), ("Achats à risque (cash insuffisant)", len(risque))]
    sections = [
        _section("Règlements à venir par jour", _lire(
            engine, f"SELECT i.date_reglement, COUNT(*) AS nb FROM instructions i WHERE {ouvert} "
                    "GROUP BY i.date_reglement ORDER BY i.date_reglement", **p), "courbe"),
        _section("Volumes par devise", _lire(
            engine, f"SELECT i.devise, COUNT(*) AS nb_instructions, SUM(i.montant) AS montant_total "
                    f"FROM instructions i WHERE {ouvert} GROUP BY i.devise ORDER BY i.devise", **p)),
        _section("Alertes : achats dont le solde cash du client est insuffisant", risque),
        _section("Détail des règlements à venir", _lire(
            engine, "SELECT i.id_instruction, c.nom AS client, t.nom AS titre, i.sens, i.quantite, i.montant, "
                    "i.devise, i.date_reglement, i.statut FROM instructions i "
                    "JOIN comptes_clients cc ON cc.id_compte = i.id_compte "
                    "JOIN clients c ON c.id_client = cc.id_client JOIN titres t ON t.isin = i.isin "
                    f"WHERE {ouvert} ORDER BY i.date_reglement, i.id_instruction", **p)),
    ]
    return indicateurs, sections


def _rapport_cash(engine, d1, d2):
    p = dict(d1=d1, d2=d2)
    entre = "m.date_valeur BETWEEN :d1 AND :d2"
    statuts = _lire(engine, f"SELECT m.statut, COUNT(*) AS nb FROM mouvements_cash m WHERE {entre} "
                            "GROUP BY m.statut", **p)
    nb = {r.statut: int(r.nb) for r in statuts.itertuples()}
    indicateurs = [("Mouvements", sum(nb.values())), ("Exécutés", nb.get("execute", 0)),
                   ("Prévus", nb.get("prevu", 0))]
    net = "SUM(CASE WHEN m.sens = 'entrant' THEN m.montant ELSE -m.montant END)"
    sections = [
        _section("Flux net prévisionnel par devise (mouvements prévus)", _lire(
            engine, f"SELECT m.devise, {net} AS flux_net FROM mouvements_cash m "
                    f"WHERE m.statut = 'prevu' AND {entre} GROUP BY m.devise ORDER BY m.devise", **p), "histogramme"),
        _section("Flux net par jour", _lire(
            engine, f"SELECT m.date_valeur, {net} AS flux_net FROM mouvements_cash m WHERE {entre} "
                    "GROUP BY m.date_valeur ORDER BY m.date_valeur", **p), "courbe"),
        _section("Entrées et sorties par devise", _lire(
            engine, f"SELECT m.devise, m.sens, COUNT(*) AS nb, SUM(m.montant) AS total FROM mouvements_cash m "
                    f"WHERE {entre} GROUP BY m.devise, m.sens ORDER BY m.devise, m.sens", **p)),
        _section("Détail des mouvements", _lire(
            engine, "SELECT m.id_mouvement, c.nom AS client, m.sens, m.montant, m.devise, m.date_valeur, m.statut "
                    "FROM mouvements_cash m JOIN comptes_clients cc ON cc.id_compte = m.id_compte "
                    f"JOIN clients c ON c.id_client = cc.id_client WHERE {entre} "
                    "ORDER BY m.date_valeur, m.id_mouvement", **p)),
    ]
    return indicateurs, sections


def _rapport_comptes(engine, d1, d2):
    nb_clients = int(_lire(engine, "SELECT COUNT(*) AS n FROM clients").iloc[0, 0])
    nb_comptes = int(_lire(engine, "SELECT COUNT(*) AS n FROM comptes_clients").iloc[0, 0])
    indicateurs = [("Clients", nb_clients), ("Comptes", nb_comptes)]
    sections = [
        _section("Clients par pays", _lire(
            engine, "SELECT pays, COUNT(*) AS nb_clients FROM clients GROUP BY pays ORDER BY nb_clients DESC"),
            "histogramme"),
        _section("Soldes par devise et type de compte", _lire(
            engine, "SELECT devise, type_compte, COUNT(*) AS nb_comptes, SUM(solde) AS solde_total "
                    "FROM comptes_clients GROUP BY devise, type_compte ORDER BY devise, type_compte")),
        _section("Soldes par client", _lire(
            engine, "SELECT c.nom AS client, cc.devise, SUM(cc.solde) AS solde_total FROM comptes_clients cc "
                    "JOIN clients c ON c.id_client = cc.id_client GROUP BY c.nom, cc.devise "
                    "ORDER BY solde_total DESC")),
    ]
    return indicateurs, sections

def _rapport_clients(engine, d1, d2):
    """Rapport dédié aux clients : nb, pays, types, liste."""
    nb_clients = int(_lire(engine, "SELECT COUNT(*) AS n FROM clients").iloc[0, 0])
    nb_pays    = int(_lire(engine, "SELECT COUNT(DISTINCT pays) AS n FROM clients").iloc[0, 0])
    nb_types   = int(_lire(engine, "SELECT COUNT(DISTINCT type_client) AS n FROM clients").iloc[0, 0])
    nb_comptes = int(_lire(engine, "SELECT COUNT(*) AS n FROM comptes_clients").iloc[0, 0])

    indicateurs = [
        ("Clients", nb_clients),
        ("Pays", nb_pays),
        ("Types de clients", nb_types),
        ("Comptes associés", nb_comptes),
    ]

    sections = [
        _section("Clients par pays", _lire(
            engine, "SELECT pays, COUNT(*) AS nb_clients FROM clients "
                    "GROUP BY pays ORDER BY nb_clients DESC"), "histogramme"),
        _section("Clients par type", _lire(
            engine, "SELECT type_client, COUNT(*) AS nb_clients FROM clients "
                    "GROUP BY type_client ORDER BY nb_clients DESC"), "histogramme"),
        _section("Liste des clients", _lire(
            engine, "SELECT c.id_client, c.nom, c.pays, c.type_client, "
                    "COUNT(cc.id_compte) AS nb_comptes "
                    "FROM clients c LEFT JOIN comptes_clients cc ON cc.id_client = c.id_client "
                    "GROUP BY c.id_client, c.nom, c.pays, c.type_client "
                    "ORDER BY c.nom")),
    ]
    return indicateurs, sections

RAPPORTS = {
    "Règlement-livraison": {"fn": _rapport_reglement, "debut": -30, "fin": 0,
                            "titre": "Rapport de règlement-livraison"},
    "Règlements à venir": {"fn": _rapport_a_venir, "debut": 0, "fin": 14,
                           "titre": "Rapport des règlements à venir"},
    "Cash et trésorerie": {"fn": _rapport_cash, "debut": -30, "fin": 14,
                           "titre": "Rapport cash et trésorerie"},
    "Comptes et soldes": {"fn": _rapport_comptes, "debut": None, "fin": None,
                          "titre": "Rapport comptes et soldes"},
    "Clients": {"fn": _rapport_clients, "debut": None, "fin": None,
                "titre": "Rapport clients"},
}

def construire_rapport(engine, nom, d1=None, d2=None):
    cfg = RAPPORTS[nom]
    indicateurs, sections = cfg["fn"](engine, d1, d2)
    periode = (f"Période du {d1:%d/%m/%Y} au {d2:%d/%m/%Y}" if cfg["debut"] is not None
               else f"Situation au {date.today():%d/%m/%Y}")
    return {"titre": cfg["titre"], "periode": periode, "genere_le": datetime.now().strftime("%d/%m/%Y %H:%M"),
            "indicateurs": indicateurs, "sections": sections, "synthese": ""}


# ---------------------------------------------------------------- export Excel
def _nom_feuille(titre, deja):
    nom = re.sub(r"[\[\]:*?/\\]", " ", titre).strip()[:28] or "Feuille"
    base, i = nom, 2
    while nom in deja:
        nom = f"{base[:25]}_{i}"
        i += 1
    deja.add(nom)
    return nom


def _mettre_en_forme(classeur):
    from openpyxl.styles import Font, PatternFill
    for ws in classeur.worksheets:
        for col in ws.columns:
            largeur = max(len(_texte(c.value)) for c in col) + 2
            ws.column_dimensions[col[0].column_letter].width = min(max(largeur, 10), 60)
        if ws.title != "Résumé":
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=ROUGE.lstrip("#"))
            ws.freeze_panes = "A2"


def exporter_excel(rap):
    """Classeur Excel : une feuille Résumé + une feuille par section."""
    lignes = [("Rapport", rap["titre"]), ("Période", rap["periode"]), ("Généré le", rap["genere_le"]), ("", "")]
    lignes += [(libelle, valeur if isinstance(valeur, (int, float)) else str(valeur))
               for libelle, valeur in rap["indicateurs"]]
    if rap.get("synthese"):
        lignes += [("", ""), ("Synthèse (Agent Rapporteur)", "")]
        lignes += [("", ligne) for ligne in rap["synthese"].splitlines() if ligne.strip()]
    tampon, deja = io.BytesIO(), {"Résumé"}
    with pd.ExcelWriter(tampon, engine="openpyxl") as xw:
        pd.DataFrame(lignes).to_excel(xw, sheet_name="Résumé", index=False, header=False)
        for s in rap["sections"]:
            s["df"].to_excel(xw, sheet_name=_nom_feuille(s["titre"], deja), index=False)
        _mettre_en_forme(xw.book)
    return tampon.getvalue()


def excel_simple(df, titre="Résultat"):
    tampon = io.BytesIO()
    with pd.ExcelWriter(tampon, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name=_nom_feuille(titre, set()), index=False)
        _mettre_en_forme(xw.book)
    return tampon.getvalue()


# ---------------------------------------------------------------- export PDF
def _graphique(df, vue, largeur):
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors

    if vue not in ("histogramme", "courbe") or len(df.columns) < 2 or len(df) < 2:
        return None
    num = [c for c in df.columns if _est_nombre(df, c)]
    x = next((c for c in df.columns if c not in num), df.columns[0])
    y = next((c for c in num if c != x), None)
    if y is None:
        return None
    d = df[[x, y]].copy()
    if vue == "histogramme":
        d = d.sort_values(y, ascending=False).head(12)
    etiquettes = [_texte(v, court=(vue == "courbe"))[:22] for v in d[x]]
    valeurs = [float(v) for v in d[y]]
    haut = 200
    dessin = Drawing(largeur, haut)
    if vue == "histogramme":
        g = VerticalBarChart()
        g.bars[0].fillColor = colors.HexColor(ROUGE)
        g.barLabelFormat = "%.0f"
        g.barLabels.nudge = 7
        g.barLabels.fontSize = 7
    else:
        g = HorizontalLineChart()
        g.lines[0].strokeColor = colors.HexColor("#4C78A8")
        g.lines[0].strokeWidth = 1.8
        pas = max(1, len(etiquettes) // 12)
        etiquettes = [e if i % pas == 0 else "" for i, e in enumerate(etiquettes)]
    g.data = [valeurs]
    g.x, g.y, g.width, g.height = 45, 55, largeur - 70, haut - 80
    g.categoryAxis.categoryNames = etiquettes
    g.categoryAxis.labels.angle = 30
    g.categoryAxis.labels.boxAnchor = "ne"
    g.categoryAxis.labels.fontSize = 7
    g.valueAxis.valueMin = min(0, min(valeurs))
    g.valueAxis.labels.fontSize = 7
    dessin.add(g)
    return dessin


def _tableau_pdf(df, largeur, max_lignes=40):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    cols = list(df.columns)
    nombres = [c for c in cols if _est_nombre(df, c)]
    a_formater = [c for c in nombres if not str(c).lower().startswith("id")]
    larg_col = largeur / len(cols)
    maxc = max(6, int(larg_col / 4.2))

    def cellule(c, v):
        s = fmt(v) if (c in a_formater and v is not None and v == v) else _texte(v)
        return s if len(s) <= maxc else s[:maxc - 1] + "."

    lignes = [[str(c)[:maxc] for c in cols]]
    for _, ligne in df.head(max_lignes).iterrows():
        lignes.append([cellule(c, ligne[c]) for c in cols])
    t = Table(lignes, colWidths=[larg_col] * len(cols), repeatRows=1)
    style = [("FONTSIZE", (0, 0), (-1, -1), 7.5), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(ROUGE)),
             ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
             ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CCCCCC")),
             ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F6F6")]),
             ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
    for i, c in enumerate(cols):
        if c in nombres:
            style.append(("ALIGN", (i, 0), (i, -1), "RIGHT"))
    t.setStyle(TableStyle(style))
    return t


def exporter_pdf(rap):
    """PDF A4 paysage : titre, indicateurs, synthèse, puis graphiques et tableaux de chaque section."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    page = landscape(A4)
    marge = 15 * mm
    largeur = page[0] - 2 * marge
    styles = getSampleStyleSheet()
    petit = ParagraphStyle("petit", parent=styles["Normal"], fontSize=8, textColor=colors.grey)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], textColor=colors.HexColor(ROUGE), spaceBefore=10)

    def pied(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(marge, 8 * mm, "Rapport post-trade - prototype, donnees de test")
        canvas.drawRightString(page[0] - marge, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    histoire = [Paragraph(escape(rap["titre"]), styles["Title"]),
                Paragraph(escape(f"{rap['periode']}  |  genere le {rap['genere_le']}"), petit), Spacer(1, 8)]

    ind = rap["indicateurs"]
    kpi = Table([[Paragraph(f"<font size=8 color='grey'>{escape(str(l))}</font>", styles["Normal"]) for l, _ in ind],
                 [Paragraph(f"<font size=15><b>{escape(fmt(v))}</b></font>", styles["Normal"]) for _, v in ind]],
                colWidths=[largeur / len(ind)] * len(ind))
    kpi.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
                             ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E0E0E0")),
                             ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FAFAFA")),
                             ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    histoire += [kpi, Spacer(1, 8)]

    if rap.get("synthese"):
        histoire.append(Paragraph("Synthese (Agent Rapporteur)", h2))
        histoire.append(Paragraph(escape(rap["synthese"]).replace("\n", "<br/>"), styles["Normal"]))

    for s in rap["sections"]:
        df = s["df"]
        entete = [Paragraph(escape(s["titre"]), h2)]
        if df.empty:
            histoire.append(KeepTogether(entete + [Paragraph("Aucune donnee pour cette section.", petit)]))
            continue
        graphique = _graphique(df, s["vue"], largeur)
        histoire.append(KeepTogether(entete + ([graphique] if graphique is not None else [])))
        histoire.append(_tableau_pdf(df, largeur))
        if len(df) > 40:
            histoire.append(Paragraph(f"... {len(df) - 40} ligne(s) supplementaire(s) : voir l'export Excel.", petit))

    tampon = io.BytesIO()
    SimpleDocTemplate(tampon, pagesize=page, leftMargin=marge, rightMargin=marge, topMargin=marge,
                      bottomMargin=marge, title=rap["titre"]).build(histoire, onFirstPage=pied, onLaterPages=pied)
    return tampon.getvalue()
